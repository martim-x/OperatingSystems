#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import io
import os
import re
import struct
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

BASE = os.path.dirname(os.path.abspath(__file__))
SORT = os.path.join(BASE, ".сортировка")
SCRY = os.path.join(BASE, "скрины")
MD = os.path.join(BASE, "вставки-для-отчета.md")
DOC = os.path.join(BASE, "lab1.docx")

TEXT_STYLE = "Для текста"
CAP_STYLE = "Normal (Web)"
LINE_FIX = {"u14-внимание-это-u13.png": "u13.png",
            "u28-внимание-это-u27.png": "u27.png"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def png_size(path):
    with open(path, "rb") as f:
        return struct.unpack(">II", f.read(24)[16:24])


def width_cm(path):
    w, _ = png_size(path)
    return 16.0 if w >= 2000 else (12.0 if w >= 1300 else 9.0)


def find_par(doc, frag, ident=""):
    hits = [p for p in doc.paragraphs if frag in p.text]
    if len(hits) != 1:
        raise RuntimeError(f"{ident}: expected 1 match for {frag!r}, got {len(hits)}")
    return hits[0]


def all_pars(doc):
    return list(doc.paragraphs)


def index_of(pars, p):
    return next(i for i, q in enumerate(pars) if q._p is p._p)


def clear_para(p):
    for node in p._p.findall(W + "r"):
        p._p.remove(node)


def set_text(p, text, style_name=TEXT_STYLE, align=None):
    clear_para(p)
    p.text = text
    if style_name:
        p.style = p.part.document.styles[style_name]
    if align is not None:
        p.alignment = align
    return p


def put_image(p, imgpath, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_para(p)
    p.alignment = align
    p.add_run().add_picture(imgpath, width=Cm(width_cm(imgpath)))
    return p


def insert_after(ref):
    new_p = OxmlElement("w:p")
    ref._p.addnext(new_p)
    return Paragraph(new_p, ref._parent)


def delete_par(p):
    p._p.getparent().remove(p._p)


def add_block_after(ref, units, doc):
    cur = ref
    for kind, data in units:
        cur = insert_after(cur)
        if kind == "text":
            set_text(cur, data, TEXT_STYLE)
        elif kind == "cap":
            set_text(cur, data, CAP_STYLE, WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "img":
            put_image(cur, data)
        else:
            set_text(cur, "")
    return cur


def parse_chapter3(md_text):
    sections = []
    cur = None
    pending_img = None
    for raw in md_text.splitlines():
        b = raw.strip()
        if not b:
            continue
        m = re.match(r"^###\s+3\.(\d+)\s+(.+)", b)
        if m:
            name = re.sub(r"\s*[–—\-]\s*рис\..*$", "", m.group(2)).strip()
            cur = [int(m.group(1)), name, []]
            sections.append(cur)
            pending_img = None
            continue
        if cur is None:
            continue
        if b.startswith("#"):
            break
        if b.startswith("(Скриншот"):
            mm = re.search(r"файл:\s*([\w.\-]+)", b)
            pending_img = mm.group(1) if mm else None
        elif b.startswith("**Рисунок 3."):
            img = pending_img
            pending_img = None
            if img:
                cur[2].append(("img", os.path.join(SORT, LINE_FIX.get(img, img))))
            cur[2].append(("cap", b.strip(" *").strip()))
        else:
            cur[2].append(("text", b))
    return sections